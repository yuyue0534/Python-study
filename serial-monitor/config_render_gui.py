# -*- coding: utf-8 -*-
"""
配置驱动 PyQt 上位机界面 Demo
--------------------------------
功能范围：
1. 读取本地 Excel 配置文件；
2. Excel 多 Sheet 自动渲染为多个 Tab 页面；
3. 每个 Sheet 按 N 行 * 4 列排布参数单元格；
4. 支持 INPUT / SELECT / SWITCH 三类控件；
5. 支持 READONLY / READWRITE；
6. 预留 Modbus RTU 读写接口：read_register / write_register；
7. 暂不接真实 Modbus，可勾选“模拟轮询值”查看页面刷新效果。

依赖安装：
    pip install PyQt5 openpyxl

运行方式：
    python config_render_gui.py
或：
    python config_render_gui.py xxx-config-example.xlsx
"""

import ast
import json
import random
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from PyQt5.QtCore import Qt, QTimer, pyqtSignal
from PyQt5.QtGui import QFont
from PyQt5.QtWidgets import (
    QApplication,
    QCheckBox,
    QComboBox,
    QFileDialog,
    QFrame,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QSpinBox,
    QStatusBar,
    QTabWidget,
    QVBoxLayout,
    QWidget,
)


@dataclass
class ConfigItem:
    name: str
    address: str
    decimals: int = 0
    unit: str = ""
    access: str = "READONLY"
    widget: str = "INPUT"
    options: List[Dict[str, Any]] = field(default_factory=list)
    sheet_name: str = ""
    row_index: int = 0

    @property
    def readonly(self) -> bool:
        return self.access.upper() == "READONLY"

    @property
    def readwrite(self) -> bool:
        return self.access.upper() == "READWRITE"


def normalize_text(value: Any) -> str:
    """把 Excel 单元格值转成干净字符串。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_int(value: Any, default: int = 0) -> int:
    try:
        if value is None or value == "":
            return default
        return int(float(value))
    except Exception:
        return default


def quote_bare_keys(text: str) -> str:
    """
    兼容类似：
        [{label:'x1', value:'v1'}]
    转成：
        [{'label':'x1', 'value':'v1'}]
    """
    return re.sub(r"([{\s,])([A-Za-z_][A-Za-z0-9_]*)\s*:", r"\1'\2':", text)


def parse_options(value: Any) -> List[Dict[str, Any]]:
    """
    支持以下几种写法：
    1. 标准 JSON：
       [{"label":"自动","value":0},{"label":"手动","value":1}]
    2. Python 字面量：
       [{'label':'自动','value':0}]
    3. 用户示例这种 key 未加引号的写法：
       [{label:'x1', value:'v1'}]
    """
    if value is None:
        return []

    if isinstance(value, list):
        raw_options = value
    else:
        text = normalize_text(value)
        if not text:
            return []

        text = (
            text.replace("，", ",")
            .replace("：", ":")
            .replace("“", '"')
            .replace("”", '"')
            .replace("‘", "'")
            .replace("’", "'")
        )

        raw_options = None

        # 先按 JSON 解析
        try:
            raw_options = json.loads(text)
        except Exception:
            pass

        # 再按 Python 字面量解析，并兼容未加引号的 label/value key
        if raw_options is None:
            try:
                raw_options = ast.literal_eval(quote_bare_keys(text))
            except Exception:
                raw_options = None

        # 最后兜底：从每个 {...} 里用正则抓 label/value
        if raw_options is None:
            raw_options = []
            for block in re.findall(r"\{(.*?)\}", text):
                label_match = re.search(r"label\s*:\s*['\"]?([^,'\"\}]+)", block)
                value_match = re.search(r"value\s*:\s*['\"]?([^,'\"\}]+)", block)
                if label_match and value_match:
                    raw_options.append(
                        {
                            "label": label_match.group(1).strip(),
                            "value": value_match.group(1).strip(),
                        }
                    )

    options: List[Dict[str, Any]] = []
    for opt in raw_options:
        if isinstance(opt, dict):
            label = opt.get("label", opt.get("name", opt.get("text", "")))
            val = opt.get("value", opt.get("val", label))
        else:
            label = str(opt)
            val = opt
        options.append({"label": str(label), "value": val})

    return options


def same_value(a: Any, b: Any) -> bool:
    """用于 SELECT 读值后匹配 options.value。"""
    if a == b:
        return True
    try:
        if float(a) == float(b):
            return True
    except Exception:
        pass
    return str(a).strip() == str(b).strip()


def read_config_excel(file_path: Path) -> Dict[str, List[ConfigItem]]:
    """
    读取 Excel 多 Sheet。
    推荐表头：
        name, address, decimals, unit, access, widget, options

    也兼容无表头：
        第1列 name
        第2列 address
        第3列 decimals
        第4列 unit
        第5列 access
        第6列 widget
        第7列 options
    """
    if not file_path.exists():
        raise FileNotFoundError(f"配置文件不存在：{file_path}")

    workbook = load_workbook(file_path, data_only=True)
    result: Dict[str, List[ConfigItem]] = {}

    default_columns = {
        "name": 0,
        "address": 1,
        "decimals": 2,
        "unit": 3,
        "access": 4,
        "widget": 5,
        "options": 6,
    }

    header_alias = {
        "name": {"name", "名称", "参数名", "标题"},
        "address": {"address", "addr", "register", "寄存器", "寄存器地址", "地址"},
        "decimals": {"decimals", "decimal", "scale", "point", "小数位", "小数点位"},
        "unit": {"unit", "单位"},
        "access": {"access", "permission", "读写权限", "权限"},
        "widget": {"widget", "type", "control", "组件", "控件类型", "类型"},
        "options": {"options", "items", "select_options", "选项", "下拉选项"},
    }

    for sheet in workbook.worksheets:
        items: List[ConfigItem] = []
        rows = list(sheet.iter_rows(values_only=True))

        # 跳过全空行，保留原始行号
        non_empty_rows = [
            (idx + 1, row)
            for idx, row in enumerate(rows)
            if any(cell is not None and str(cell).strip() != "" for cell in row)
        ]

        if not non_empty_rows:
            result[sheet.title] = []
            continue

        first_row_number, first_row = non_empty_rows[0]
        lowered_first = [normalize_text(cell).strip().lower() for cell in first_row]

        # 判断是否有表头
        has_header = "name" in lowered_first and (
            "address" in lowered_first or "addr" in lowered_first or "寄存器地址" in lowered_first
        )

        if has_header:
            columns = {}
            for col_idx, title in enumerate(lowered_first):
                for standard_key, aliases in header_alias.items():
                    if title in aliases:
                        columns[standard_key] = col_idx
                        break
            start_data_pos = 1
        else:
            columns = default_columns
            start_data_pos = 0

        def get_value(row_data: tuple, key: str, default: Any = "") -> Any:
            col = columns.get(key)
            if col is None or col >= len(row_data):
                return default
            value = row_data[col]
            return default if value is None else value

        for actual_row_number, row in non_empty_rows[start_data_pos:]:
            name = normalize_text(get_value(row, "name"))
            address = normalize_text(get_value(row, "address"))

            # name/address 为空时跳过
            if not name or not address:
                continue

            decimals = normalize_int(get_value(row, "decimals"), default=0)
            unit = normalize_text(get_value(row, "unit"))
            access = normalize_text(get_value(row, "access", "READONLY")).upper() or "READONLY"
            widget = normalize_text(get_value(row, "widget", "INPUT")).upper() or "INPUT"
            options = parse_options(get_value(row, "options", ""))

            # 容错：SELECT 没有 options 时给一个空选项，避免界面空白
            if widget == "SELECT" and not options:
                options = [{"label": "未配置选项", "value": ""}]

            items.append(
                ConfigItem(
                    name=name,
                    address=address,
                    decimals=decimals,
                    unit=unit,
                    access=access,
                    widget=widget,
                    options=options,
                    sheet_name=sheet.title,
                    row_index=actual_row_number,
                )
            )

        result[sheet.title] = items

    return result


class ParamCell(QFrame):
    """单个参数单元格：左边 label，右边 INPUT / SELECT / SWITCH，再加单位。"""

    write_requested = pyqtSignal(object, object)  # ConfigItem, raw_value

    def __init__(self, item: ConfigItem, parent: Optional[QWidget] = None):
        super().__init__(parent)
        self.item = item
        self._updating = False

        self.setObjectName("ParamCell")
        self.setFrameShape(QFrame.StyledPanel)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)

        self.title_label = QLabel(item.name)
        self.title_label.setMinimumWidth(86)
        self.title_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.title_label.setToolTip(
            f"Sheet: {item.sheet_name}\n"
            f"Excel 行号: {item.row_index}\n"
            f"寄存器地址: {item.address}\n"
            f"小数位: {item.decimals}\n"
            f"权限: {item.access}\n"
            f"控件: {item.widget}"
        )

        self.unit_label = QLabel(item.unit)
        self.unit_label.setMinimumWidth(26)
        self.unit_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)

        self.editor = self._create_editor()

        layout = QHBoxLayout(self)
        layout.setContentsMargins(8, 6, 8, 6)
        layout.setSpacing(6)
        layout.addWidget(self.title_label)
        layout.addWidget(self.editor, 1)
        layout.addWidget(self.unit_label)

        self._apply_style()

    def _create_editor(self) -> QWidget:
        widget_type = self.item.widget.upper()

        if widget_type == "SELECT":
            combo = QComboBox()
            for opt in self.item.options:
                combo.addItem(str(opt["label"]), opt["value"])
            combo.setEnabled(self.item.readwrite)
            combo.activated.connect(self._on_select_activated)
            return combo

        if widget_type in {"SWITCH", "CHECKBOX", "BOOL", "BOOLEAN"}:
            checkbox = QCheckBox("开启")
            checkbox.setEnabled(self.item.readwrite)
            checkbox.stateChanged.connect(self._on_switch_changed)
            return checkbox

        # 默认 INPUT
        line_edit = QLineEdit()
        line_edit.setPlaceholderText("等待读值")
        line_edit.setReadOnly(self.item.readonly)
        if self.item.readwrite:
            line_edit.editingFinished.connect(self._on_input_finished)
        return line_edit

    def _apply_style(self) -> None:
        base = """
        QFrame#ParamCell {
            border: 1px solid #D0D7DE;
            border-radius: 8px;
            background: #FFFFFF;
        }
        QLabel {
            color: #24292F;
        }
        QLineEdit, QComboBox {
            min-height: 26px;
            border: 1px solid #D0D7DE;
            border-radius: 5px;
            padding: 2px 6px;
            background: #FFFFFF;
        }
        QLineEdit[readOnly="true"], QComboBox:disabled, QCheckBox:disabled {
            color: #57606A;
            background: #F6F8FA;
        }
        """
        self.setStyleSheet(base)

    def _on_input_finished(self) -> None:
        if self._updating or not isinstance(self.editor, QLineEdit):
            return

        text = self.editor.text().strip()
        if text == "":
            return

        try:
            # 显示值 -> 原始寄存器值
            raw_value = int(round(float(text) * (10 ** self.item.decimals)))
        except ValueError:
            QMessageBox.warning(self, "输入错误", f"“{self.item.name}”请输入数字。")
            return

        self.write_requested.emit(self.item, raw_value)

    def _on_select_activated(self, index: int) -> None:
        if self._updating or not isinstance(self.editor, QComboBox):
            return

        raw_value = self.editor.itemData(index)
        self.write_requested.emit(self.item, raw_value)

    def _on_switch_changed(self, state: int) -> None:
        if self._updating:
            return
        raw_value = 1 if state == Qt.Checked else 0
        self.write_requested.emit(self.item, raw_value)

    def set_display_value(self, raw_value: Any) -> None:
        """
        轮询读到的寄存器原始值 -> 控件显示。
        INPUT：按 decimals 缩放；
        SELECT：匹配 options.value；
        SWITCH：0/1 显示关/开。
        """
        self._updating = True
        try:
            widget_type = self.item.widget.upper()

            if widget_type == "SELECT" and isinstance(self.editor, QComboBox):
                matched_index = -1
                for i in range(self.editor.count()):
                    if same_value(self.editor.itemData(i), raw_value):
                        matched_index = i
                        break
                if matched_index >= 0:
                    self.editor.setCurrentIndex(matched_index)
                else:
                    # 读到未配置值时，不强行新增选项，只在 tooltip 提醒
                    self.editor.setToolTip(f"读到未配置值：{raw_value}")
                return

            if widget_type in {"SWITCH", "CHECKBOX", "BOOL", "BOOLEAN"} and isinstance(self.editor, QCheckBox):
                self.editor.setChecked(str(raw_value).strip() not in {"0", "False", "false", ""})
                return

            if isinstance(self.editor, QLineEdit):
                try:
                    scaled = float(raw_value) / (10 ** self.item.decimals)
                    if self.item.decimals <= 0:
                        display_text = str(int(round(scaled)))
                    else:
                        display_text = f"{scaled:.{self.item.decimals}f}"
                except Exception:
                    display_text = str(raw_value)
                self.editor.setText(display_text)
        finally:
            self._updating = False


class MainWindow(QMainWindow):
    def __init__(self, config_path: Optional[Path] = None):
        super().__init__()

        self.setWindowTitle("配置驱动上位机 GUI Demo（Modbus 预留版）")
        self.resize(1280, 760)

        self.config_path = config_path
        self.config_data: Dict[str, List[ConfigItem]] = {}
        self.param_cells: List[ParamCell] = []
        self.polling = False

        self._build_ui()
        self._build_menu()

        if config_path and config_path.exists():
            self.load_config(config_path)
        else:
            self.statusBar().showMessage("未找到默认配置文件，请通过 文件 -> 打开配置文件 加载 Excel。")

    def _build_menu(self) -> None:
        file_menu = self.menuBar().addMenu("文件")
        open_action = file_menu.addAction("打开配置文件...")
        open_action.triggered.connect(self.choose_config_file)

        reload_action = file_menu.addAction("重新加载当前配置")
        reload_action.triggered.connect(self.reload_current_config)

    def _build_ui(self) -> None:
        central = QWidget()
        main_layout = QVBoxLayout(central)
        main_layout.setContentsMargins(12, 10, 12, 10)
        main_layout.setSpacing(10)

        main_layout.addWidget(self._create_serial_area())

        self.tab_widget = QTabWidget()
        self.tab_widget.setDocumentMode(True)
        main_layout.addWidget(self.tab_widget, 1)

        self.setCentralWidget(central)
        self.setStatusBar(QStatusBar())

        self.setStyleSheet("""
        QMainWindow {
            background: #F6F8FA;
        }
        QGroupBox {
            font-weight: bold;
            border: 1px solid #D0D7DE;
            border-radius: 8px;
            margin-top: 10px;
            background: #FFFFFF;
        }
        QGroupBox::title {
            subcontrol-origin: margin;
            left: 12px;
            padding: 0 4px;
        }
        QPushButton {
            min-height: 28px;
            border-radius: 6px;
            padding: 4px 12px;
            border: 1px solid #D0D7DE;
            background: #FFFFFF;
        }
        QPushButton:hover {
            background: #F6F8FA;
        }
        QTabWidget::pane {
            border: 1px solid #D0D7DE;
            background: #FFFFFF;
        }
        QTabBar::tab {
            padding: 8px 18px;
            border: 1px solid #D0D7DE;
            background: #F6F8FA;
            margin-right: 2px;
        }
        QTabBar::tab:selected {
            background: #FFFFFF;
            font-weight: bold;
        }
        """)

    def _create_serial_area(self) -> QGroupBox:
        group = QGroupBox("RTU 串口配置（预留）")
        layout = QHBoxLayout(group)
        layout.setContentsMargins(12, 18, 12, 10)
        layout.setSpacing(8)

        self.port_combo = QComboBox()
        self.baud_combo = QComboBox()
        self.data_bits_combo = QComboBox()
        self.stop_bits_combo = QComboBox()
        self.parity_combo = QComboBox()
        self.poll_interval_spin = QSpinBox()
        self.simulate_checkbox = QCheckBox("模拟轮询值")
        self.refresh_port_btn = QPushButton("刷新串口")
        self.poll_btn = QPushButton("开始轮询预览")

        self.baud_combo.addItems(["9600", "19200", "38400", "57600", "115200"])
        self.baud_combo.setCurrentText("9600")
        self.data_bits_combo.addItems(["8", "7"])
        self.stop_bits_combo.addItems(["1", "1.5", "2"])
        self.parity_combo.addItems(["None", "Even", "Odd"])

        self.poll_interval_spin.setRange(0, 60000)
        self.poll_interval_spin.setValue(100)
        self.poll_interval_spin.setSuffix(" ms")

        self.refresh_port_btn.clicked.connect(self.refresh_ports)
        self.poll_btn.clicked.connect(self.toggle_polling)

        layout.addWidget(QLabel("串口"))
        layout.addWidget(self.port_combo)
        layout.addWidget(self.refresh_port_btn)

        layout.addSpacing(10)
        layout.addWidget(QLabel("波特率"))
        layout.addWidget(self.baud_combo)

        layout.addWidget(QLabel("数据位"))
        layout.addWidget(self.data_bits_combo)

        layout.addWidget(QLabel("停止位"))
        layout.addWidget(self.stop_bits_combo)

        layout.addWidget(QLabel("校验"))
        layout.addWidget(self.parity_combo)

        layout.addWidget(QLabel("轮询间隔"))
        layout.addWidget(self.poll_interval_spin)

        layout.addWidget(self.simulate_checkbox)
        layout.addStretch(1)
        layout.addWidget(self.poll_btn)

        self.refresh_ports()
        return group

    def refresh_ports(self) -> None:
        ports = self.list_serial_ports()
        self.port_combo.clear()
        self.port_combo.addItems(ports)
        self.statusBar().showMessage(f"已刷新串口列表：{len(ports)} 个候选项。", 3000)

    @staticmethod
    def list_serial_ports() -> List[str]:
        """
        不强制依赖 pyserial。
        如果已安装 pyserial，则读取真实串口；否则给 Windows 常用 COM1-COM20 作为候选。
        """
        try:
            from serial.tools import list_ports  # type: ignore

            ports = [port.device for port in list_ports.comports()]
            if ports:
                return ports
        except Exception:
            pass

        if sys.platform.startswith("win"):
            return [f"COM{i}" for i in range(1, 21)]
        return ["/dev/ttyUSB0", "/dev/ttyUSB1", "/dev/ttyS0", "/dev/ttyS1"]

    def choose_config_file(self) -> None:
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "选择配置文件",
            str(Path.cwd()),
            "Excel 配置文件 (*.xlsx *.xlsm);;所有文件 (*.*)",
        )
        if file_path:
            self.load_config(Path(file_path))

    def reload_current_config(self) -> None:
        if not self.config_path:
            QMessageBox.information(self, "提示", "当前没有已加载的配置文件。")
            return
        self.load_config(self.config_path)

    def load_config(self, file_path: Path) -> None:
        try:
            self.config_data = read_config_excel(file_path)
        except Exception as exc:
            QMessageBox.critical(self, "配置读取失败", str(exc))
            return

        self.config_path = file_path
        self.render_tabs()
        self.statusBar().showMessage(f"配置加载成功：{file_path}", 5000)

    def render_tabs(self) -> None:
        self.tab_widget.clear()
        self.param_cells.clear()

        for sheet_name, items in self.config_data.items():
            page = QWidget()
            grid = QGridLayout(page)
            grid.setContentsMargins(14, 14, 14, 14)
            grid.setHorizontalSpacing(10)
            grid.setVerticalSpacing(10)

            if not items:
                empty_label = QLabel("该 Sheet 没有可渲染的配置项。")
                empty_label.setAlignment(Qt.AlignCenter)
                grid.addWidget(empty_label, 0, 0)
            else:
                for idx, item in enumerate(items):
                    row = idx // 4
                    col = idx % 4

                    cell = ParamCell(item)
                    cell.write_requested.connect(self.write_register)
                    self.param_cells.append(cell)
                    grid.addWidget(cell, row, col)

                # 4 列等宽拉伸
                for col in range(4):
                    grid.setColumnStretch(col, 1)

            scroll = QScrollArea()
            scroll.setWidgetResizable(True)
            scroll.setWidget(page)
            self.tab_widget.addTab(scroll, sheet_name)

        self.statusBar().showMessage(
            f"已渲染 {len(self.config_data)} 个 Tab，共 {len(self.param_cells)} 个配置项。", 5000
        )

    def toggle_polling(self) -> None:
        self.polling = not self.polling
        if self.polling:
            self.poll_btn.setText("停止轮询预览")
            self.statusBar().showMessage("轮询框架已启动。当前没有真实 Modbus；勾选“模拟轮询值”可查看刷新效果。")
            self.poll_once()
        else:
            self.poll_btn.setText("开始轮询预览")
            self.statusBar().showMessage("轮询已停止。", 3000)

    def poll_once(self) -> None:
        """
        轮询周期逻辑：
        1. 顺序读取本周期全部寄存器；
        2. 本周期结束后，再等待 poll_interval_spin 设置的毫秒数；
        3. 如果间隔为 0，则不额外等待，下一轮进入事件队列后立即执行。

        真实 Modbus 接入时，建议把 read_register 改成异步或线程池执行，避免阻塞 UI。
        """
        if not self.polling:
            return

        for cell in self.param_cells:
            raw_value = self.read_register(cell.item)
            if raw_value is not None:
                cell.set_display_value(raw_value)

        interval_ms = self.poll_interval_spin.value()
        QTimer.singleShot(interval_ms, self.poll_once)

    def read_register(self, item: ConfigItem) -> Optional[Any]:
        """
        预留：Modbus 读寄存器入口。

        你后续接 pymodbus / minimalmodbus 时，可以在这里按 item.address 读取值：
            raw_value = client.read_holding_registers(...)
            return raw_value

        当前 Demo 默认不返回值。
        勾选“模拟轮询值”后，返回随机值，用来观察页面刷新效果。
        """
        if not self.simulate_checkbox.isChecked():
            return None

        widget_type = item.widget.upper()

        if widget_type == "SELECT" and item.options:
            return random.choice(item.options)["value"]

        if widget_type in {"SWITCH", "CHECKBOX", "BOOL", "BOOLEAN"}:
            return random.choice([0, 1])

        # INPUT：模拟原始寄存器值；显示时会按 decimals 缩放
        max_display_value = 100
        return random.randint(0, max_display_value * (10 ** item.decimals))

    def write_register(self, item: ConfigItem, raw_value: Any) -> None:
        """
        预留：Modbus 写寄存器入口。

        INPUT：
            用户输入的是显示值，比如 decimals=2，输入 12.34；
            这里收到的 raw_value 是 1234。

        SELECT：
            用户选中的是 options 里的 value 字段；
            这里收到的 raw_value 就是那个 value。

        SWITCH：
            关闭=0，开启=1。
        """
        # TODO：后续在这里接真实 Modbus 写寄存器逻辑
        # 示例：
        # client.write_register(address=int(item.address), value=int(raw_value), slave=1)

        msg = f"写入预留：{item.name} | 地址={item.address} | 原始值={raw_value}"
        print(msg)
        self.statusBar().showMessage(msg, 5000)


def resolve_default_config_path() -> Optional[Path]:
    """
    优先级：
    1. 命令行第一个参数；
    2. 当前 py 文件同目录下的 xxx-config-example.xlsx；
    3. 当前工作目录下的 xxx-config-example.xlsx。
    """
    if len(sys.argv) >= 2:
        return Path(sys.argv[1]).expanduser().resolve()

    script_dir_path = Path(__file__).resolve().with_name("xxx-config-example.xlsx")
    if script_dir_path.exists():
        return script_dir_path

    cwd_path = Path.cwd() / "xxx-config-example.xlsx"
    if cwd_path.exists():
        return cwd_path

    return None


def main() -> None:
    app = QApplication(sys.argv)

    # 中文界面字体设置：Windows 优先微软雅黑，其他系统没有也不会报错
    font = QFont("Microsoft YaHei", 9)
    app.setFont(font)

    config_path = resolve_default_config_path()
    window = MainWindow(config_path=config_path)
    window.show()

    sys.exit(app.exec_())


if __name__ == "__main__":
    main()
